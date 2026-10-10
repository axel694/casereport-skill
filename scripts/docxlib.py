"""Hjälpbibliotek för rapport (docx): Times New Roman 12, rubriker, punktlistor, tabeller och bilder.
Användning: from docxlib import *   (sätt docxlib.IMG till mappen med PNG-diagram innan IMGP anropas)
Skapa diagram med matplotlib (Times New Roman) och spara som PNG i IMG. fmt/pct ger svensk talformatering."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
IMG = os.path.join(os.getcwd(), "diagram")
def fmt(v, d=0, sign=False):
    s = f"{abs(v):,.{d}f}".replace(",", " ").replace(".", ",")
    return ("−" + s) if v < 0 else (("+" + s) if sign else s)
def pct(v, d=0, sign=False): return fmt(v * 100, d, sign) + " %"
# ---------- dokument
doc = Document()
sec = doc.sections[0]; sec.left_margin = sec.right_margin = Cm(2.3); sec.top_margin = sec.bottom_margin = Cm(2.2)
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(6)

def run_fmt(r, bold=False, italic=False, size=12, color=None):
    r.font.name = "Times New Roman"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman"); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = RGBColor.from_string(color)
def H(t):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14); p.paragraph_format.keep_with_next = True; run_fmt(p.add_run(t), bold=True)
def H2(t):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.keep_with_next = True; run_fmt(p.add_run(t), bold=True, italic=True)
def Pp(t, bold=False, italic=False, size=12):
    p = doc.add_paragraph(); run_fmt(p.add_run(t), bold=bold, italic=italic, size=size); return p
def Bl(items):
    for t in items:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2); run_fmt(p.add_run(t))
def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)
def T(data, widths=None, size=10, bold_rows=(), hl=None):
    t = doc.add_table(rows=len(data), cols=len(data[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(data):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ""; p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(v)); run_fmt(r, bold=(i == 0 or i in bold_rows or (hl and (i, j) in hl)), size=size, color="FFFFFF" if i == 0 else None)
            if j > 0: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i == 0: shade(c, "1F3864")
            elif hl and (i, j) in hl: shade(c, "EFE9E1")
            elif i % 2 == 0: shade(c, "F2F2F2")
            if widths: c.width = Cm(widths[j])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
def IMGP(name, w=15.5):
    doc.add_picture(f"{IMG}/{name}.png", width=Cm(w)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
def note(t): Pp(t, italic=True, size=10)

