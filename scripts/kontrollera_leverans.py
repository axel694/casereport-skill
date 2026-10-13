#!/usr/bin/env python3
"""Leveranskontroll: python3 kontrollera_leverans.py DECK.pptx RAPPORT.docx
1) Punktdecimaler (21.367) i svensk text, 2) shapes utanför sliden eller överlappande, 3) täckning: tal i slidetext (utom diagram) som saknas i rapporten,
4) förbjudna ord (licensierade terminaler). Tal i diagram ska ligga som tabell i rapporten (bilaga med diagramunderlag)."""
import re, sys
from pptx import Presentation
from docx import Document
pp, dd = Presentation(sys.argv[1]), Document(sys.argv[2])
dt = " ".join([p.text for p in dd.paragraphs] + [c.text for t in dd.tables for r in t.rows for c in r.cells])
norm = lambda x: x.replace("\xa0", " ").replace(" ", "")
pt = []
for sl in pp.slides:
    for sh in sl.shapes:
        if sh.has_text_frame: pt.append(sh.text_frame.text)
        if sh.has_table: pt += [c.text for r in sh.table.rows for c in r.cells]
    pt.append(sl.notes_slide.notes_text_frame.text)
alltxt = " ".join(pt) + " " + dt
dots = sorted(set(re.findall(r"(?<![\d.])\d+\.\d+(?![\d.])", alltxt)))
print("Punktdecimaler (kontrollera, versionsnummer/exhibit är ok):", dots)
W, H = pp.slide_width, pp.slide_height; prob = 0
for i, sl in enumerate(pp.slides, 1):
    sh = [s for s in sl.shapes if s.width and s.height]
    for a in range(len(sh)):
        s = sh[a]
        if s.left < 0 or s.top < 0 or s.left + s.width > W + 10 or s.top + s.height > H + 10: print(f"slide {i}: utanför sliden: {s.name}"); prob += 1
        for b in range(a + 1, len(sh)):
            t = sh[b]; ox = min(s.left + s.width, t.left + t.width) - max(s.left, t.left); oy = min(s.top + s.height, t.top + t.height) - max(s.top, t.top)
            if ox > 20000 and oy > 20000 and not (s.name.startswith("Rectangle") or t.name.startswith("Rectangle")): print(f"slide {i}: överlapp {s.name} / {t.name}"); prob += 1
dn = norm(dt); miss = []
for i, sl in enumerate(pp.slides, 1):
    for sh in sl.shapes:
        txts = ([sh.text_frame.text] if sh.has_text_frame else []) + ([c.text for r in sh.table.rows for c in r.cells] if sh.has_table else [])
        for t in txts:
            for tok in re.findall(r"[−+]?\d[\d  ]*(?:,\d+)?\s?(?:%|x|mn|mdr)?", t):
                if len(re.sub(r"\D", "", tok)) >= 2 and norm(tok) not in dn: miss.append((i, tok.strip()))
print(f"Layoutproblem: {prob}. Tal på slides som saknas i rapporten ({len(miss)}; formatskillnader som '5,4 mdr' mot '5 362 mn' är ok): {miss}")
print("Förbjudna ord:", [w for w in ["Bloomberg", "FactSet", "Refinitiv", "Capital IQ"] if w.lower() in alltxt.lower()])
