"""Hjälpfunktioner för slides: textrutor, tabeller, inbyggda diagram och en enkel passningskontroll."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from PIL import ImageFont

NAVY, SAND, RED, GREY, WHITE, BLACK = "1F3864", "B8A38D", "C00000", "D9D9D9", "FFFFFF", "000000"
LGREY, LRED, LNAVY, LSAND = "F2F2F2", "F8E3E3", "DCE3F0", "EFE9E1"
SANS, SERIF = "Calibri", "Times New Roman"
FIT = []          # (slide_no, beskrivning, beräknad höjd, boxhöjd)

try:
    _F = {False: ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 100),
          True: ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 100)}
except Exception:
    _F = None

def rgb(h): return RGBColor.from_string(h)

def fmt(v, d=0, sign=False):
    s = f"{abs(v):,.{d}f}".replace(",", " ").replace(".", ",")
    if v < 0: return "−" + s
    return ("+" + s) if sign else s

def pct(v, d=0, sign=False): return fmt(v * 100, d, sign) + " %"

def est_lines(text, size, width_in, bold=False):
    """Uppskattar antal rader (Arial-mått × 0.9 ≈ Calibri). Ger marginal åt säkra sidan."""
    if _F is None: return max(1, len(text) * size * 0.5 / 72 / width_in)
    wpx = width_in * 72 / size * 100 / 0.9       # bredd i "100-punkts-pixlar" justerad för Calibri
    lines = 0
    for para in text.split("\n"):
        words, cur, n = para.split(" "), "", 1
        for w in words:
            t = (cur + " " + w).strip()
            if _F[bold].getlength(t) > wpx and cur: n += 1; cur = w
            else: cur = t
        lines += n
    return lines

class Deck:
    def __init__(self):
        self.p = Presentation(); self.p.slide_width, self.p.slide_height = Inches(13.333), Inches(7.5)
        self.n = 0
    def slide(self, title, subtitle=None, source=None):
        s = self.p.slides.add_slide(self.p.slide_layouts[6]); self.n += 1
        s._no = self.n
        bar = s.shapes.add_shape(1, 0, 0, self.p.slide_width, Inches(0.08)); bar.fill.solid(); bar.fill.fore_color.rgb = rgb(NAVY); bar.line.fill.background()
        self.text(s, 0.4, 0.15, 12.5, 0.55, [(title, 24, True, NAVY, SERIF)], name="titel")
        if subtitle: self.text(s, 0.4, 0.68, 12.5, 0.32, [(subtitle, 12, False, "595959", SANS)], name="undertitel")
        if source: self.text(s, 0.4, 7.08, 12.5, 0.32, [(source, 8, False, "595959", SANS)], name="källa")
        return s
    def text(self, s, x, y, w, h, paras, fill=None, line=None, name="ruta", anchor="top", pad=0.06, bullet=False, check=True):
        """paras: lista av (text, storlek, fet, färg, typsnitt[, bullet]) eller str."""
        sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        if line: sh.line.color.rgb = rgb(line); sh.line.width = Pt(0.75)
        tf = sh.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(pad); tf.margin_top = tf.margin_bottom = Inches(0.04)
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}[anchor]
        tot = 0.08
        for i, p in enumerate(paras):
            if isinstance(p, str): p = (p, 11, False, BLACK, SANS)
            t, sz, b, col, fnt = (list(p) + [None] * 5)[:5]
            bl = bullet and len(p) > 5 and p[5]
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = para.add_run(); r.text = ("• " + t) if bl else t
            r.font.size = Pt(sz); r.font.bold = bool(b); r.font.name = fnt or SANS; r.font.color.rgb = rgb(col or BLACK)
            para.space_after = Pt(2)
            tot += est_lines(r.text, sz, w - 2 * pad - 0.05, bool(b)) * sz * 1.2 / 72 + 2 / 72
        if check: FIT.append((s._no, name, round(tot, 2), h))
        return sh
    def bullets(self, s, x, y, w, h, items, size=11, head=None, headcol=NAVY, fill=None, line=None, name="punkter"):
        paras = []
        if head: paras.append((head, size + 1, True, headcol, SANS))
        for t in items: paras.append((t, size, False, BLACK, SANS, True))
        return self.text(s, x, y, w, h, paras, fill=fill, line=line, bullet=True, name=name)
    def table(self, s, x, y, w, data, colw=None, size=9, rowh=0.26, head_fill=NAVY, bold_first_col=False, hl=None, red_rows=(), bold_rows=(), name="tabell", align_first="left"):
        rows, cols = len(data), len(data[0])
        gs = s.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(rowh * rows))
        tbl = gs.table
        tblPr = tbl._tbl.tblPr
        tblPr.set("firstRow", "1"); tblPr.set("bandRow", "0")
        sid = tblPr.find(qn("a:tableStyleId"))
        if sid is not None: sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # ingen stil, inga linjer
        if colw:
            tot = sum(colw)
            for i, c in enumerate(colw): tbl.columns[i].width = Inches(w * c / tot)
        est_h = 0
        for r in range(rows):
            maxl = 1
            for c in range(cols):
                cell = tbl.cell(r, c); v = str(data[r][c]) or "\u00a0"
                cell.margin_left = cell.margin_right = Inches(0.05); cell.margin_top = cell.margin_bottom = Inches(0.02)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame; tf.word_wrap = True
                p = tf.paragraphs[0]; run = p.add_run(); run.text = v
                run.font.size = Pt(size); run.font.name = SANS
                head = r == 0
                run.font.bold = head or (bold_first_col and c == 0) or (r in bold_rows)
                run.font.color.rgb = rgb(WHITE if head else (RED if r in red_rows else BLACK))
                p.alignment = PP_ALIGN.LEFT if (c == 0 and align_first == "left") else (PP_ALIGN.LEFT if head and c == 0 else PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(head_fill if head else (LGREY if r % 2 == 0 else WHITE))
                if hl and (r, c) in hl: cell.fill.fore_color.rgb = rgb(LSAND); run.font.bold = True
                cw = (w * (colw[c] / sum(colw)) if colw else w / cols) - 0.1
                maxl = max(maxl, est_lines(v, size, cw, run.font.bold))
            h = max(rowh, maxl * size * 1.2 / 72 + 0.06)
            tbl.rows[r].height = Inches(h); est_h += h
        FIT.append((s._no, name + " (höjd)", round(est_h, 2), None))
        return gs, est_h
    def chart(self, s, x, y, w, h, cats, series, kind="col", colors=(NAVY, SAND, RED), numfmt='#,##0', title=None, legend=True, labels=True, size=9, gap=60, point_colors=None, ymax=None):
        cd = CategoryChartData(); cd.categories = cats
        for n, vals in series: cd.add_series(n, vals)
        ct = {"col": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "bar": XL_CHART_TYPE.BAR_CLUSTERED}[kind]
        gf = s.shapes.add_chart(ct, Inches(x), Inches(y), Inches(w), Inches(h), cd); ch = gf.chart
        ch.font.size = Pt(size); ch.font.name = SANS
        ch.has_legend = legend and len(series) > 1
        if ch.has_legend: ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(size)
        if title:
            ch.has_title = True; ch.chart_title.text_frame.text = title
            ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(size + 2); ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
        else: ch.has_title = False
        va = ch.value_axis; va.has_major_gridlines = True; va.major_gridlines.format.line.color.rgb = rgb("E0E0E0")
        va.tick_labels.font.size = Pt(size); va.tick_labels.number_format = numfmt; va.tick_labels.number_format_is_linked = False
        ch.category_axis.tick_labels.font.size = Pt(size)
        if ymax is not None: va.maximum_scale = ymax; va.minimum_scale = 0
        for i, ser in enumerate(ch.plots[0].series):
            col = colors[i % len(colors)]
            if kind == "line":
                ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(2.25); ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(col); ser.smooth = False
            else:
                ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(col)
        if kind != "line": ch.plots[0].gap_width = gap
        if point_colors and kind != "line":
            ser = ch.plots[0].series[0]
            for idx, col in enumerate(point_colors):
                pt = ser.points[idx]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(col)
        if labels:
            pl = ch.plots[0]; pl.has_data_labels = True; dl = pl.data_labels
            dl.font.size = Pt(size - 1); dl.number_format = numfmt; dl.number_format_is_linked = False
            if kind != "line": dl.position = XL_LABEL_POSITION.OUTSIDE_END
            else: dl.position = XL_LABEL_POSITION.ABOVE
        return gf
    def report_fit(self):
        bad = [f for f in FIT if f[3] is not None and f[2] > f[3] + 0.02]
        for f in bad: print("FÖR HÖG:", f)
        print(f"{len(FIT)} kontroller, {len(bad)} för höga")
        return bad
