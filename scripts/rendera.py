import subprocess, shutil, os, sys
from pptx import Presentation
src = sys.argv[1]; only = [int(x) for x in sys.argv[2:]]
shutil.rmtree("/tmp/qs", ignore_errors=True); os.makedirs("/tmp/qs"); os.makedirs("/tmp/ql", exist_ok=True)
n = len(Presentation(src).slides)
for k in (only or range(1, n + 1)):
    p = Presentation(src); ids = p.slides._sldIdLst
    for i, sid in enumerate(list(ids)):
        if i != k - 1: p.part.drop_rel(sid.rId); ids.remove(sid)
    f = f"/tmp/qs/s{k:02d}.pptx"; p.save(f)
    subprocess.run(["qlmanage", "-t", "-s", "1600", "-o", "/tmp/ql", f], capture_output=True)
